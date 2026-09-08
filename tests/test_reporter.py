import pytest
from pathlib import Path
from decimal import Decimal
from datetime import datetime

import openpyxl

from foxprice.core.reporter import Reporter
from foxprice.core.models import PriceOffer, PricedResult, ErrorResult, RunSummary


@pytest.fixture
def summary() -> RunSummary:
    return RunSummary(
        run_id="test0001",
        parts_total=3,
        parts_found=2,
        parts_not_found=1,
        errors=1,
        duration_sec=5.5,
        suppliers=["sup_a"],
        mer_score={"sup_a": 1},
    )


@pytest.fixture
def offer() -> PriceOffer:
    return PriceOffer(
        supplier="sup_a",
        part_number="P-001",
        matched_part="P-001",
        unit_price=Decimal("5.00"),
        warehouse_label="NYC",
    )


@pytest.fixture
def priced(offer) -> list[PricedResult]:
    return [
        PricedResult(offer=offer, markup_pct=Decimal("200"), final_price=Decimal("15.00"))
    ]


@pytest.fixture
def errors() -> list[ErrorResult]:
    return [
        ErrorResult(
            supplier="sup_a",
            part_number="P-002",
            error_type="not_found",
            description="no match",
            attempt=2,
        )
    ]


def test_generate_creates_file(tmp_path, priced, errors, summary):
    Reporter(tmp_path).generate(priced, errors, summary)
    assert (tmp_path / "foxprice_test0001.xlsx").exists()


def test_workbook_has_four_sheets(tmp_path, priced, errors, summary):
    path = Reporter(tmp_path).generate(priced, errors, summary)
    wb = openpyxl.load_workbook(path)
    assert set(wb.sheetnames) == {"Results", "Raw Data", "Errors", "Settings"}


def test_results_sheet_data(tmp_path, priced, errors, summary):
    path = Reporter(tmp_path).generate(priced, errors, summary)
    wb = openpyxl.load_workbook(path)
    ws = wb["Results"]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][0] == "Supplier"
    assert rows[1][0] == "sup_a"
    assert rows[1][3] == "5.00"
    assert rows[1][5] == "15.00"


def test_errors_sheet_data(tmp_path, priced, errors, summary):
    path = Reporter(tmp_path).generate(priced, errors, summary)
    wb = openpyxl.load_workbook(path)
    ws = wb["Errors"]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[1][2] == "not_found"
    assert rows[1][4] == 2


def test_settings_sheet_has_run_id(tmp_path, priced, errors, summary):
    path = Reporter(tmp_path).generate(priced, errors, summary)
    wb = openpyxl.load_workbook(path)
    ws = wb["Settings"]
    keys = [row[0] for row in ws.iter_rows(values_only=True)]
    assert "Run ID" in keys


def test_empty_priced_no_crash(tmp_path, errors, summary):
    Reporter(tmp_path).generate([], errors, summary)


def test_output_dir_created(tmp_path):
    subdir = tmp_path / "nested" / "output"
    Reporter(subdir).generate(
        [],
        [],
        RunSummary(
            run_id="x",
            parts_total=0,
            parts_found=0,
            parts_not_found=0,
            errors=0,
            duration_sec=0.0,
            suppliers=[],
            mer_score={},
        ),
    )
    assert subdir.exists()
